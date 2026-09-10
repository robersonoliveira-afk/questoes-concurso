# -*- coding: utf-8 -*-
"""Gera a planilha de revisão manual das questões com problema.
Lê rev_tmp.json (produzido pelo scan) e escreve revisao_questoes.xlsx
na pasta do projeto Teste_site_questões (que o OneDrive sincroniza
entre as duas máquinas do Róberson).
"""
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

linhas = json.load(open("rev_tmp.json", encoding="utf-8"))
linhas.sort(key=lambda r: (r["ano"], r["numero"]))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Revisar"

COLS = [
    ("id", 16, False),
    ("ano", 6, False),
    ("nº", 5, False),
    ("disciplina", 16, False),
    ("por que revisar", 24, False),
    ("bloco", 10, True),
    ("TEXTO-BASE", 55, True),
    ("ENUNCIADO", 55, True),
    ("A", 30, True),
    ("B", 30, True),
    ("C", 30, True),
    ("D", 30, True),
    ("E", 30, True),
    ("gabarito", 9, True),
    ("figura?", 8, True),
    ("(referência) o que está no sistema hoje", 60, False),
]

azul = PatternFill("solid", fgColor="DCE6F1")     # colunas que ele preenche
cinza = PatternFill("solid", fgColor="EFEFEF")    # referência / não mexer
hdrfill = PatternFill("solid", fgColor="1F3864")
hdrfont = Font(color="FFFFFF", bold=True, size=10)

for j, (nome, larg, editavel) in enumerate(COLS, 1):
    c = ws.cell(row=1, column=j, value=nome)
    c.fill = hdrfill
    c.font = hdrfont
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width = larg

for i, r in enumerate(linhas, 2):
    vals = [
        r["id"], r["ano"], r["numero"], r["disciplina"], r["motivo"],
        f'{r["ano"]}-{r["disciplina"][:3]}',   # palpite de bloco, ele ajusta
        "",                                     # TEXTO-BASE
        "",                                     # ENUNCIADO
        "", "", "", "", "",                     # A-E
        r["gabarito"],                          # gabarito pré-preenchido
        r["tem_figura"],                        # figura?
        r["enunciado_atual"][:1200],
    ]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=j, value=v)
        cell.alignment = Alignment(vertical="top", wrap_text=(COLS[j - 1][2]))
        if COLS[j - 1][0] in ("TEXTO-BASE", "ENUNCIADO", "A", "B", "C", "D", "E", "gabarito", "figura?", "bloco"):
            cell.fill = azul
        elif COLS[j - 1][0].startswith("(referência)"):
            cell.fill = cinza
    ws.row_dimensions[i].height = 90

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# ---- aba de instruções ----
ins = wb.create_sheet("Como preencher")
texto = [
    ("Como preencher esta planilha", True),
    ("", False),
    ("Preencha só as colunas AZUIS. Não mexa em id, ano, nº.", False),
    ("", False),
    ("TEXTO-BASE", True),
    ("Se a questão diz 'leia o texto para responder às questões X a Y', cole o texto aqui.", False),
    ("Preencha só na PRIMEIRA questão do bloco e deixe as outras em branco — eu repito na importação.", False),
    ("Se a questão não tem texto compartilhado, deixe em branco.", False),
    ("", False),
    ("ENUNCIADO", True),
    ("O comando da questão, limpo, sem cabeçalho de página nem número de linha.", False),
    ("", False),
    ("A, B, C, D, E", True),
    ("As cinco alternativas, uma por coluna. Só o texto, sem a letra.", False),
    ("", False),
    ("Fórmulas (Matemática, Química)", True),
    ("Escreva entre cifrões. Exemplos:", False),
    ("  $x^2$        ->  x²", False),
    ("  $\\sqrt{21}$   ->  raiz de 21", False),
    ("  $\\frac{1}{2}$ ->  um meio", False),
    ("  $2\\pi$        ->  2 pi", False),
    ("  $2x^3 + 4x^2 - \\frac{3}{4} = 0$", False),
    ("Se preferir não mexer com isso, escreva do jeito que der e eu ajusto depois.", False),
    ("", False),
    ("figura?", True),
    ("Marque 's' se a questão TEM uma figura essencial (mapa, gráfico, diagrama, tirinha, tabela).", False),
    ("Eu recorto a figura do PDF — você não precisa fazer isso.", False),
    ("", False),
    ("gabarito", True),
    ("Confira a letra. Já vem preenchida com o que o sistema tem hoje (quase sempre certo).", False),
    ("", False),
    ("Não precisa fazer tudo de uma vez. Salve o arquivo quando terminar um lote e me avise.", False),
]
for i, (t, bold) in enumerate(texto, 1):
    c = ins.cell(row=i, column=1, value=t)
    if bold:
        c.font = Font(bold=True, size=11)
ins.column_dimensions["A"].width = 100

destino = Path(__file__).resolve().parents[2] / "Teste_site_questões" / "revisao_questoes.xlsx"
wb.save(destino)
print(f"{len(linhas)} questões -> {destino}")
